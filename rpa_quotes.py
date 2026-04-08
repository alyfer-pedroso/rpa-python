"""
╔══════════════════════════════════════════════════════════════╗
║         RPA - Automação de Tarefas Repetitivas               ║
║         Site REAL: http://quotes.toscrape.com                ║
║         Saída: Planilha Excel formatada + JSON               ║
╚══════════════════════════════════════════════════════════════╝

INSTALAÇÃO:
    pip install -r requirements.txt

USO:
    python rpa_quotes.py                → coleta tudo
    python rpa_quotes.py love           → filtra pela tag "love"
    python rpa_quotes.py inspirational  → filtra pela tag "inspirational"
"""

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
import json
import sys
import time
from datetime import datetime
from collections import Counter

CONFIG = {
    "url_base":            "http://quotes.toscrape.com",
    "delay_entre_paginas": 1.0,
    "timeout":             10,
    "max_paginas":         100,
    "filtro_tag":          None,
    "salvar_json":         True,
    "arquivo_json":        "relatorio_rpa.json",
    "arquivo_xlsx":        "relatorio_rpa.xlsx",
}

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                  "AppleWebKit/537.36 (KHTML, like Gecko) "
                  "Chrome/120.0.0.0 Safari/537.36"
}

COR = {
    "azul_escuro":  "1A3A5C",
    "azul_medio":   "2E6DA4",
    "azul_claro":   "D6E4F0",
    "branco":       "FFFFFF",
    "cinza_claro":  "F4F6F9",
    "cinza_medio":  "DDE3EA",
    "texto_escuro": "1C2B3A",
    "texto_medio":  "4A5E72",
    "verde_bg":     "E8F5E9",
    "verde_fg":     "2E7D32",
}


def cabecalho():
    print("\n" + "═" * 62)
    print("  🤖  RPA - Automação de Coleta de Dados")
    print(f"  🌐  Site: {CONFIG['url_base']}")
    print(f"  📅  Início: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
    if CONFIG["filtro_tag"]:
        print(f"  🏷   Filtro por tag: '{CONFIG['filtro_tag']}'")
    print("═" * 62 + "\n")


def acessar_pagina(url: str):
    """Faz requisição HTTP real e retorna BeautifulSoup."""
    try:
        print(f"  🌐 [GET] {url}")
        resp = requests.get(url, headers=HEADERS, timeout=CONFIG["timeout"])
        resp.raise_for_status()
        print(f"       → {resp.status_code} OK  ({len(resp.content)} bytes)")
        return BeautifulSoup(resp.text, "html.parser")
    except requests.exceptions.ConnectionError:
        print("  ❌ Sem conexão com a internet.")
        return None
    except requests.exceptions.Timeout:
        print(f"  ❌ Timeout ao acessar {url}")
        return None
    except requests.exceptions.HTTPError as e:
        print(f"  ❌ Erro HTTP: {e}")
        return None


def extrair_citacoes(soup) -> list:
    citacoes = []
    for bloco in soup.select("div.quote"):
        texto = bloco.select_one("span.text").get_text(strip=True)
        autor = bloco.select_one("small.author").get_text(strip=True)
        tags  = [t.get_text(strip=True) for t in bloco.select("a.tag")]
        citacoes.append({
            "texto":    texto,
            "autor":    autor,
            "tags":     tags,
            "tags_str": ", ".join(tags),
        })
    return citacoes


def proxima_pagina_url(soup) -> str:
    btn = soup.select_one("li.next > a")
    return CONFIG["url_base"] + btn["href"] if btn else None


def filtrar_por_tag(citacoes: list, tag: str) -> list:
    return [c for c in citacoes if tag.lower() in [t.lower() for t in c["tags"]]]


def coletar_dados() -> list:
    cabecalho()
    todas = []
    url_atual = CONFIG["url_base"]
    pagina = 1

    print("  🔄 Iniciando navegação automática...\n")
    while url_atual and pagina <= CONFIG["max_paginas"]:
        print(f"  📄 Página {pagina}:")
        soup = acessar_pagina(url_atual)
        if soup is None:
            print("  ⚠️  Encerrando por erro de conexão.")
            break
        lote = extrair_citacoes(soup)
        todas.extend(lote)
        print(f"       ✔ {len(lote)} citações extraídas.\n")
        url_atual = proxima_pagina_url(soup)
        pagina += 1
        if url_atual:
            time.sleep(CONFIG["delay_entre_paginas"])

    print(f"  🏁 Navegação concluída! {len(todas)} citações em {pagina-1} página(s).")

    if CONFIG["filtro_tag"]:
        print(f"\n  🔍 Filtrando por tag: '{CONFIG['filtro_tag']}'...")
        todas = filtrar_por_tag(todas, CONFIG["filtro_tag"])
        print(f"     ✔ {len(todas)} citações após filtro.")

    return todas


def salvar_json(citacoes: list):
    dados = {
        "gerado_em": datetime.now().isoformat(),
        "site":      CONFIG["url_base"],
        "total":     len(citacoes),
        "citacoes":  citacoes,
    }
    with open(CONFIG["arquivo_json"], "w", encoding="utf-8") as f:
        json.dump(dados, f, ensure_ascii=False, indent=2)
    print(f"  💾 JSON salvo → {CONFIG['arquivo_json']}")

def borda(estilo="thin", cor="BBCAD6"):
    s = Side(style=estilo, color=cor)
    return Border(left=s, right=s, top=s, bottom=s)

def fill(cor):
    return PatternFill("solid", fgColor=cor)

def fonte(tamanho=10, negrito=False, cor="1C2B3A", italico=False):
    return Font(name="Arial", size=tamanho, bold=negrito,
                color=cor, italic=italico)

def alinhar(h="left", v="center", wrap=False, indent=0):
    return Alignment(horizontal=h, vertical=v,
                     wrap_text=wrap, indent=indent)

def aplicar(cell, value=None, font=None, bg=None,
            align=None, border=None, height=None, row=None, ws=None):
    if value is not None:
        cell.value = value
    if font:
        cell.font = font
    if bg:
        cell.fill = fill(bg)
    if align:
        cell.alignment = align
    if border:
        cell.border = border
    if height and row and ws:
        ws.row_dimensions[row].height = height

def construir_aba_citacoes(wb: Workbook, citacoes: list):
    ws = wb.active
    ws.title = "📋 Citações"

    ws.merge_cells("A1:D1")
    aplicar(ws["A1"],
            value="🤖  RPA — Citações Coletadas  |  quotes.toscrape.com",
            font=fonte(14, negrito=True, cor=COR["branco"]),
            bg=COR["azul_escuro"],
            align=alinhar("center", "center"))
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A2:D2")
    aplicar(ws["A2"],
            value=f"Total de registros: {len(citacoes)}   •   Gerado automaticamente pelo RPA",
            font=fonte(10, italico=True, cor=COR["texto_medio"]),
            bg=COR["azul_claro"],
            align=alinhar("center", "center"))
    ws.row_dimensions[2].height = 20

    ws.row_dimensions[3].height = 6 

    colunas = [("A", "#", 5), ("B", "Citação", 72),
               ("C", "Autor", 22), ("D", "Tags", 38)]
    for col, label, largura in colunas:
        c = ws[f"{col}4"]
        aplicar(c,
                value=label,
                font=fonte(10, negrito=True, cor=COR["branco"]),
                bg=COR["azul_medio"],
                align=alinhar("center", "center", wrap=True),
                border=borda())
        ws.column_dimensions[col].width = largura
    ws.row_dimensions[4].height = 24

    for i, row in enumerate(citacoes, 1):
        r = i + 4
        bg = COR["cinza_claro"] if i % 2 == 0 else COR["branco"]

        aplicar(ws.cell(r, 1, value=i),
                font=fonte(10, negrito=True, cor=COR["azul_medio"]),
                bg=bg, align=alinhar("center", "top"), border=borda())

        aplicar(ws.cell(r, 2, value=row["texto"]),
                font=fonte(10),
                bg=bg, align=alinhar("left", "top", wrap=True), border=borda())

        aplicar(ws.cell(r, 3, value=row["autor"]),
                font=fonte(10, negrito=True, cor=COR["azul_escuro"]),
                bg=bg, align=alinhar("center", "top", wrap=True), border=borda())

        aplicar(ws.cell(r, 4, value=row["tags_str"]),
                font=fonte(9, italico=True, cor=COR["texto_medio"]),
                bg=bg, align=alinhar("left", "top", wrap=True), border=borda())

        ws.row_dimensions[r].height = 48

    total_row = len(citacoes) + 5
    ws.merge_cells(f"A{total_row}:D{total_row}")
    aplicar(ws[f"A{total_row}"],
            value=f"TOTAL DE CITAÇÕES: {len(citacoes)}",
            font=fonte(10, negrito=True, cor=COR["branco"]),
            bg=COR["azul_escuro"],
            align=alinhar("center", "center"))
    ws.row_dimensions[total_row].height = 22

    ws.freeze_panes = "A5"

def construir_aba_estatisticas(wb: Workbook, citacoes: list):
    ws = wb.create_sheet("📊 Estatísticas")

    todas_tags     = [t.strip() for c in citacoes for t in c["tags_str"].split(",") if t.strip()]
    autores_count  = Counter(c["autor"] for c in citacoes).most_common(8)
    tags_count     = Counter(todas_tags).most_common(10)
    autores_unicos = len(set(c["autor"] for c in citacoes))
    tags_unicas    = len(set(todas_tags))

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 4
    ws.column_dimensions["D"].width = 28
    ws.column_dimensions["E"].width = 14

    ws.merge_cells("A1:E1")
    aplicar(ws["A1"],
            value="📊  Estatísticas Gerais — RPA quotes.toscrape.com",
            font=fonte(14, negrito=True, cor=COR["branco"]),
            bg=COR["azul_escuro"],
            align=alinhar("center", "center"))
    ws.row_dimensions[1].height = 36

    ws.merge_cells("A3:B3")
    aplicar(ws["A3"],
            value="RESUMO GERAL",
            font=fonte(11, negrito=True, cor=COR["branco"]),
            bg=COR["azul_medio"],
            align=alinhar("center", "center"))
    ws.row_dimensions[3].height = 22

    resumo = [
        ("Total de Citações",     len(citacoes)),
        ("Autores Únicos",        autores_unicos),
        ("Tags Únicas",           tags_unicas),
        ("Média de Tags/Citação", f"=ROUND({len(todas_tags)}/{len(citacoes)},1)"),
    ]
    for i, (label, val) in enumerate(resumo):
        r = i + 4
        aplicar(ws.cell(r, 1, value=label),
                font=fonte(10, negrito=True),
                bg=COR["azul_claro"],
                align=alinhar("left", "center", indent=1),
                border=borda())
        aplicar(ws.cell(r, 2, value=val),
                font=fonte(11, negrito=True, cor=COR["azul_escuro"]),
                bg=COR["branco"],
                align=alinhar("center", "center"),
                border=borda())
        ws.row_dimensions[r].height = 22

    ws.merge_cells("A9:B9")
    aplicar(ws["A9"],
            value="TOP AUTORES",
            font=fonte(11, negrito=True, cor=COR["branco"]),
            bg=COR["azul_medio"],
            align=alinhar("center", "center"))
    ws.row_dimensions[9].height = 22

    for col, label in [(1, "Autor"), (2, "Citações")]:
        aplicar(ws.cell(10, col, value=label),
                font=fonte(10, negrito=True, cor=COR["azul_escuro"]),
                bg=COR["cinza_medio"],
                align=alinhar("center", "center"),
                border=borda())
    ws.row_dimensions[10].height = 20

    for i, (autor, cnt) in enumerate(autores_count):
        r = i + 11
        bg = COR["cinza_claro"] if i % 2 == 0 else COR["branco"]
        aplicar(ws.cell(r, 1, value=autor),
                font=fonte(10, negrito=True),
                bg=bg,
                align=alinhar("left", "center", indent=1),
                border=borda())
        aplicar(ws.cell(r, 2, value=cnt),
                font=fonte(10, negrito=True, cor=COR["azul_medio"]),
                bg=bg,
                align=alinhar("center", "center"),
                border=borda())
        ws.row_dimensions[r].height = 20

    ws.merge_cells("D3:E3")
    aplicar(ws["D3"],
            value="TOP 10 TAGS",
            font=fonte(11, negrito=True, cor=COR["branco"]),
            bg=COR["azul_medio"],
            align=alinhar("center", "center"))

    for col, label in [(4, "Tag"), (5, "Ocorrências")]:
        aplicar(ws.cell(4, col, value=label),
                font=fonte(10, negrito=True, cor=COR["azul_escuro"]),
                bg=COR["cinza_medio"],
                align=alinhar("center", "center"),
                border=borda())
    ws.row_dimensions[4].height = 20

    for i, (tag, cnt) in enumerate(tags_count):
        r = i + 5
        bg = COR["verde_bg"] if i % 2 == 0 else COR["branco"]
        aplicar(ws.cell(r, 4, value=tag),
                font=fonte(10, cor=COR["verde_fg"]),
                bg=bg,
                align=alinhar("left", "center", indent=1),
                border=borda())
        aplicar(ws.cell(r, 5, value=cnt),
                font=fonte(10, negrito=True, cor=COR["verde_fg"]),
                bg=bg,
                align=alinhar("center", "center"),
                border=borda())
        ws.row_dimensions[r].height = 20

    chart = BarChart()
    chart.type       = "col"
    chart.title      = "Citações por Autor"
    chart.y_axis.title = "Quantidade"
    chart.x_axis.title = "Autor"
    chart.style      = 10
    chart.height     = 10
    chart.width      = 18
    chart.grouping   = "clustered"

    n = len(autores_count)
    data_ref = Reference(ws, min_col=2, min_row=10, max_row=10 + n)
    cats_ref = Reference(ws, min_col=1, min_row=11, max_row=10 + n)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    chart.series[0].graphicalProperties.solidFill = COR["azul_medio"]

    ws.add_chart(chart, "A20")

def gerar_planilha(citacoes: list):
    wb = Workbook()
    construir_aba_citacoes(wb, citacoes)
    construir_aba_estatisticas(wb, citacoes)
    wb.save(CONFIG["arquivo_xlsx"])
    print(f"  📊 Excel salvo → {CONFIG['arquivo_xlsx']}")


def exibir_relatorio(citacoes: list):
    print("\n" + "═" * 62)
    print("  📊  RELATÓRIO FINAL")
    print("═" * 62)
    print(f"\n  ✅ Total de citações coletadas: {len(citacoes)}")

    autores = Counter(c["autor"] for c in citacoes)
    print("\n  👤 Top 5 Autores:")
    for autor, qtd in autores.most_common(5):
        print(f"     {'█' * qtd} {autor} ({qtd}x)")

    todas_tags = [t.strip() for c in citacoes for t in c["tags_str"].split(",") if t.strip()]
    print("\n  🏷️  Top 10 Tags:")
    for tag, qtd in Counter(todas_tags).most_common(10):
        print(f"     • {tag:35s} {qtd}x")

    print("\n  💬 Citações Coletadas:")
    for i, c in enumerate(citacoes, 1):
        trecho = c["texto"][:75] + "..." if len(c["texto"]) > 75 else c["texto"]
        print(f"\n  [{i:02d}] {trecho}")
        print(f"        — {c['autor']}")
        print(f"        🏷  {c['tags_str']}")

    print("\n" + "═" * 62)
    print(f"  ✔  Concluído em {datetime.now().strftime('%H:%M:%S')}")
    print("═" * 62 + "\n")


def executar_rpa():
    citacoes = coletar_dados()

    print("\n  💾 Salvando arquivos...")
    if CONFIG["salvar_json"]:
        salvar_json(citacoes)
    gerar_planilha(citacoes)

    exibir_relatorio(citacoes)


if __name__ == "__main__":
    if len(sys.argv) > 1:
        CONFIG["filtro_tag"] = sys.argv[1]
    executar_rpa()